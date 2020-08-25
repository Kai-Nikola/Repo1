import openpyxl as xl
from openpyxl.chart import BarChart, Reference

wb = xl.load_workbook('RETO.xlsx')

sheet = wb['Sheet1']


values = Reference(sheet, min_row=2, max_row=sheet.max_row, min_col=2, max_col=3)

chart = BarChart()
chart.add_data(values)
sheet.add_chart(chart, 'd2')

wb.save('RETO_res2.xlsx')
